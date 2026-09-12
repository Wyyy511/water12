from __future__ import annotations
from datetime import datetime
import math, re
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from core.paths import OUTPUT_DIR
from agent.risk_brief import baseline_brief, scenario_brief

CJK="Noto Sans CJK SC"

def _missing(v):
    return v is None or (isinstance(v,float) and (math.isnan(v) or math.isinf(v)))

def _fmt(v,n=4):
    if _missing(v): return "不可计算"
    try: return f"{float(v):.{n}f}"
    except Exception: return "不可计算"

def _pct(v):
    if _missing(v): return "不可计算"
    try: return f"{float(v)*100:.1f}%"
    except Exception: return "不可计算"

def _clean_md(s:str)->str:
    s=re.sub(r"\*\*(.*?)\*\*",r"\1",s)
    s=re.sub(r"^#+\s*","",s)
    return s.replace("`","").strip()

def _shade(cell, fill="EAF3FF"):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)

def _set_cell_text(cell,text,bold=False):
    cell.text=""
    r=cell.paragraphs[0].add_run(str(text))
    r.bold=bold; r.font.name=CJK
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),CJK)

def _force_cjk(doc):
    for st in doc.styles:
        try:
            st.font.name=CJK; st._element.rPr.rFonts.set(qn("w:eastAsia"),CJK)
        except Exception: pass
    for para in doc.paragraphs:
        for run in para.runs:
            run.font.name=CJK; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),CJK)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.name=CJK; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),CJK)

def _add_ai_assessment(doc,text,heading):
    doc.add_heading(heading,1)
    lines=[x.strip() for x in (text or "").splitlines() if x.strip()]
    if not lines:
        doc.add_paragraph("本次 AI 解读未完整生成。风险计算结果仍保留，可在 WaterPulse 中重试 AI 解读。")
        return
    for line in lines:
        clean=_clean_md(line)
        if not clean: continue
        if line.startswith("###") or any(clean.startswith(k) for k in [
            "风险结论","风险结果","重点供应地","主要风险来源","数据可信度与限制","数据说明","管理建议",
            "变化结果","主要变化来源"
        ]):
            p=doc.add_paragraph(); p.add_run(clean).bold=True
        elif line.startswith("- "):
            doc.add_paragraph(_clean_md(line[2:]),style="List Bullet")
        else:
            doc.add_paragraph(clean)

def _primary(b):
    data=(b or {}).get("data") or {}
    s=data.get("summary"); n=data.get("nodes")
    if isinstance(s,list): s=pd.DataFrame(s)
    if isinstance(n,list): n=pd.DataFrame(n)
    if s is None or getattr(s,"empty",True): return None,None,None
    r=s.iloc[0]
    gn=n
    if isinstance(n,pd.DataFrame) and not n.empty and "material" in n.columns:
        gn=n[n.material==r.get("material")]
    top=None
    if isinstance(gn,pd.DataFrame) and not gn.empty and "C" in gn.columns:
        top=gn.sort_values("C",ascending=False).iloc[0]
    vals={"长期缺水":float(r.get("path_contrib_WS") or 0),
          "历史干旱":float(r.get("path_contrib_DR") or 0),
          "季节波动":float(r.get("path_contrib_SV") or 0)}
    dom=max(vals,key=vals.get) if sum(vals.values())>0 else None
    return r,top,dom

def _scenario_summary(s):
    if not s or s.get("data") is None: return None
    d=s["data"]; typ=d.get("scenario_type")
    if typ=="NodeFailure":
        return {
            "压力测试":"主要供应地中断",
            "目标供应地":d.get("target_node_name") or d.get("target_node_id") or "—",
            "中断比例":_pct(d.get("failure_fraction_f")),
            "预计供应损失":_fmt(d.get("gross_loss")),
            "库存缓冲":_fmt(d.get("inventory_used")),
            "替代供应":_fmt(d.get("replacement_allocated")),
            "未满足需求":_fmt(d.get("unmet_demand")),
        }
    if typ in {"PeakSeason","ExtremeDrought"}:
        return {"压力测试":"关键用水期压力" if typ=="PeakSeason" else "严重干旱",
                "当前风险":_fmt(d.get("PRWI_baseline")),
                "测试后风险":_fmt(d.get("PRWI_scenario")),
                "变化":_fmt(d.get("PRWI_delta"))}
    if typ=="AqueductFuture":
        return {"压力测试":"未来水环境变化","年份":d.get("year"),"路径":d.get("path"),
                "当前风险":_fmt(d.get("PRWI_baseline")),
                "测试后风险":_fmt(d.get("PRWI_future")),
                "变化":_fmt(d.get("PRWI_delta"))}
    return {"压力测试":typ or "已运行"}

def export_risk_report(baseline_result:dict, scenario_result:dict|None, validation:dict|None, snapshot_id:str="",
                       baseline_explanation:str="", scenario_explanation:str="")->str:
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
    doc.styles["Normal"].font.name=CJK;doc.styles["Normal"].font.size=Pt(10.5)

    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=p.add_run("WaterPulse\n上游供应链水风险评估与管理建议报告");rr.bold=True;rr.font.size=Pt(22)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("面向企业采购、供应链、ESG 与风险管理决策").italic=True

    meta=doc.add_table(rows=3,cols=2);meta.style="Table Grid";meta.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,(k,v) in enumerate([
        ("报告生成时间",datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("分析编号",snapshot_id or (baseline_result or {}).get("run_id") or "—"),
        ("报告用途","识别上游采购水风险、测试供应链韧性并形成行动建议")
    ]):
        _set_cell_text(meta.cell(i,0),k,True);_shade(meta.cell(i,0));_set_cell_text(meta.cell(i,1),v)

    r,top,dom=_primary(baseline_result)
    proxies=(baseline_result or {}).get("proxy_fields") or []
    assumptions=(baseline_result or {}).get("assumptions") or []
    ident=(baseline_result or {}).get("data_identity") or {}

    doc.add_heading("1. 执行摘要",1)
    if r is None:
        doc.add_paragraph("当前资料不足以形成完整的企业水风险评估。请先补充采购来源、供应地区和采购占比。")
    else:
        ent=str(r.get("enterprise") or "本次分析对象");mat=str(r.get("material") or "原材料")
        doc.add_paragraph(
            f"本报告针对 {ent} 的 {mat} 上游采购组合进行水风险筛查。当前综合相对水风险指数为 {_fmt(r.get('PRWI'))}。"
            "该指标用于识别和比较风险重点，不代表断供概率，也不直接等于财务损失。"
        )
        if top is not None:
            doc.add_paragraph(f"当前最需要关注的供应地为 {top.get('node_name')}（{top.get('node_id')}），对采购组合的风险贡献占比约 {_pct(top.get('contribution_share'))}。")
        vals={"长期缺水":float(r.get("path_contrib_WS") or 0),"历史干旱":float(r.get("path_contrib_DR") or 0),"季节波动":float(r.get("path_contrib_SV") or 0)}
        sm=sum(vals.values())
        if dom and sm:
            doc.add_paragraph(f"最主要的风险来源为 {dom}，约占当前综合风险的 {vals[dom]/sm*100:.1f}%。管理建议应优先针对这一风险来源和最大贡献供应地。")
        if ident.get("has_proxy") or ident.get("has_assumption"):
            doc.add_paragraph("本次结果含已明确标记的代理数据或研究假设，因此适合用于风险筛查与采购决策准备，不应被解释为全部企业数据已完成真实核验。")
        ss=_scenario_summary(scenario_result)
        if ss and ss.get("压力测试")=="主要供应地中断":
            doc.add_paragraph(f"压力测试显示：若 {ss.get('目标供应地')} 按设定发生中断，预计供应损失 {ss.get('预计供应损失')}，经库存与替代供应缓冲后，未满足需求为 {ss.get('未满足需求')}。")

    doc.add_heading("2. 分析对象与数据基础",1)
    data=(baseline_result or {}).get("data") or {}
    s=data.get("summary")
    if isinstance(s,list): s=pd.DataFrame(s)
    if isinstance(s,pd.DataFrame) and not s.empty:
        t=doc.add_table(rows=1,cols=7);t.style="Table Grid"
        heads=["企业","原材料","综合风险","采购覆盖","有效数据覆盖","未识别份额","证据可信度"]
        for i,h in enumerate(heads):_set_cell_text(t.rows[0].cells[i],h,True);_shade(t.rows[0].cells[i],"DCEBFA")
        for _,x in s.iterrows():
            vals=[x.get("enterprise") or "未提供",x.get("material") or "未提供",_fmt(x.get("PRWI")),
                  _pct(x.get("Coverage")),_pct(x.get("scored_coverage")),_pct(x.get("unknown_share_U")),x.get("overall_confidence") or "Unknown"]
            cells=t.add_row().cells
            for i,v in enumerate(vals):_set_cell_text(cells[i],v)
    if ident:
        doc.add_paragraph("数据身份："+str(ident.get("label") or "未标记")+"。"+str(ident.get("note") or ""))
    if proxies or assumptions:
        doc.add_paragraph(f"本次分析包含代理数据 {len(proxies)} 项、假设参数 {len(assumptions)} 项。重大采购决策前应尽可能用企业真实数据替换。")
    if validation:
        notes=(validation.get("issues") or [])+(validation.get("data_gaps") or [])
        if notes: doc.add_paragraph("仍需进一步核实："+"；".join(map(str,notes[:10])))

    doc.add_heading("3. 当前上游供应链水风险",1)
    if r is not None:
        vals=[("长期缺水压力",float(r.get("path_contrib_WS") or 0)),
              ("历史干旱",float(r.get("path_contrib_DR") or 0)),
              ("季节性供水波动",float(r.get("path_contrib_SV") or 0))]
        sm=sum(v for _,v in vals)
        t=doc.add_table(rows=1,cols=3);t.style="Table Grid"
        for i,h in enumerate(["风险来源","风险贡献","占综合风险比例"]):_set_cell_text(t.rows[0].cells[i],h,True);_shade(t.rows[0].cells[i],"DCEBFA")
        for label,v in vals:
            cells=t.add_row().cells
            _set_cell_text(cells[0],label);_set_cell_text(cells[1],_fmt(v,6));_set_cell_text(cells[2],_pct(v/sm if sm else None))
        nodes=data.get("nodes")
        if isinstance(nodes,list):nodes=pd.DataFrame(nodes)
        if isinstance(nodes,pd.DataFrame) and not nodes.empty:
            if "material" in nodes.columns:nodes=nodes[nodes.material==r.get("material")]
            if "C" in nodes.columns:nodes=nodes.sort_values("C",ascending=False).head(5)
            doc.add_paragraph("风险贡献最高的供应节点：")
            nt=doc.add_table(rows=1,cols=5);nt.style="Table Grid"
            for i,h in enumerate(["排序","供应地","节点","贡献值","贡献占比"]):_set_cell_text(nt.rows[0].cells[i],h,True);_shade(nt.rows[0].cells[i],"EAF3FF")
            for rank,(_,node) in enumerate(nodes.iterrows(),1):
                vals2=[rank,node.get("node_name") or "—",node.get("node_id") or "—",_fmt(node.get("C"),6),_pct(node.get("contribution_share"))]
                cells=nt.add_row().cells
                for i,v in enumerate(vals2):_set_cell_text(cells[i],v)

    _add_ai_assessment(doc,baseline_explanation.strip() or baseline_brief(baseline_result),"4. AI 风险评估与管理建议")

    doc.add_heading("5. 压力测试与供应链韧性",1)
    if scenario_result:
        ss=_scenario_summary(scenario_result) or {}
        st=doc.add_table(rows=0,cols=2);st.style="Table Grid"
        for k,v in ss.items():
            cells=st.add_row().cells;_set_cell_text(cells[0],k,True);_shade(cells[0]);_set_cell_text(cells[1],v)
        _add_ai_assessment(doc,scenario_explanation.strip() or scenario_brief(scenario_result),"压力测试解读与应对建议")
    else:
        doc.add_paragraph("本次未运行压力测试。企业可继续运行“主要供应地中断”测试，评估关键节点中断时库存、替代供应和未满足需求的变化。")

    doc.add_heading("6. 建议企业优先采取的行动",1)
    actions=[]
    if top is not None:
        actions.append(f"优先核实并持续监测 {top.get('node_name')} 的供水与供应稳定性，因为该供应地对当前采购组合的风险贡献最高。")
    if dom=="长期缺水":
        actions.append("针对长期缺水风险，核查高水压力供应地的取水依赖、灌溉来源和替代产区可行性。")
    elif dom=="历史干旱":
        actions.append("针对干旱风险，建立关键供应地干旱预警与采购触发机制，并评估库存与替代来源是否足以覆盖阶段性减产。")
    elif dom=="季节波动":
        actions.append("针对季节性供水波动，对齐关键采购期与当地高水压力月份，必要时调整采购时点、库存或来源组合。")
    if proxies or assumptions:
        actions.append("在重大采购调整前，优先用企业真实采购、供应商和产地数据替换当前代理值或研究假设。")
    if scenario_result and (scenario_result.get("data") or {}).get("scenario_type")=="NodeFailure":
        actions.append("根据供应地中断测试结果，明确库存缓冲、备选供应商和企业可接受的最大未满足需求阈值。")
    actions.append("将本报告作为采购与供应链风险筛查依据，并结合成本、质量、合同和供应商尽调进行人工决策。")
    for x in actions:doc.add_paragraph(x,style="List Bullet")

    doc.add_heading("7. 数据来源、限制与可追溯信息",1)
    details=(baseline_result or {}).get("source_details") or []
    if details:
        t=doc.add_table(rows=1,cols=6);t.style="Table Grid"
        heads=["分析对象","来源机构/项目标识","原始来源/DOI","参考期","数据身份","置信度"]
        for i,h in enumerate(heads):_set_cell_text(t.rows[0].cells[i],h,True);_shade(t.rows[0].cells[i],"DCEBFA")
        for d in details[:25]:
            vals3=[d.get("node_name") or d.get("material") or d.get("node_id") or "数据项",
                   d.get("source_org") or d.get("secondary_source_org") or d.get("source_ref") or "待补",
                   d.get("original_url") or d.get("method_doi") or d.get("secondary_method_doi") or "待补",
                   d.get("reference_period") or "待补",d.get("data_type") or "待补",d.get("confidence") or "待补"]
            cells=t.add_row().cells
            for i,v in enumerate(vals3):_set_cell_text(cells[i],v)
    doc.add_paragraph("本报告用于上游水风险筛查、供应节点排序、采购结构比较和压力测试。综合风险指数不是断供概率或财务损失预测。代理采购组合或示范采购组合不得表述为真实企业风险敞口。")

    doc.add_heading("附录：模型与运行信息",1)
    doc.add_paragraph(f"风险计算引擎：{(baseline_result or {}).get('engine_version') or '待记录'}；数据版本：{(baseline_result or {}).get('data_version') or '待记录'}；运行编号：{(baseline_result or {}).get('run_id') or snapshot_id or '待记录'}。")
    checks=(baseline_result or {}).get("integrity_checks") or []
    if checks:
        doc.add_paragraph("计算一致性校验："+("通过" if all(x.get("ok") for x in checks) else "未通过"))
    doc.add_paragraph("AI 负责理解需求、解释确定性计算结果并形成管理建议；核心风险数值与压力测试数字由确定性程序产生，不由语言模型自行生成。")

    _force_cjk(doc)
    out=OUTPUT_DIR/f"WaterPulse_Enterprise_Water_Risk_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    doc.save(out);return str(out)
